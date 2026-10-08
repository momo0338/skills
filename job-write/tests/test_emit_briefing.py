#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""emit_briefing / workspace_source 测试（2026-09-30 数据源改造后）。

背景
====
`job-write` 原先把外部项目 ijob（`ijob.public` brief 契约）当唯一事实源。
2026-09-30 已彻底解耦：取数改为「码上职业」**本工作区自有**的三个事实源，
适配层在同目录 `scripts/workspace_source.py`，由 `scripts/emit_briefing.py` 调用：

  ① `单位与招聘信息总表.xlsx`：`招聘信息`（批次台账）+ `公司信息`（单位档案）
  ② `巡检记录/*监测.md`：在架分层（有效 / 长期 / 待核）
  ③ `内容状态.json`：写没写过 / 推没推草稿箱 / 发没发 —— 唯一真源

因此本文件：
  - **不含任何 ijob 相关测试**（无 IJOB_SKILL_SCRIPT / ijob.py / brief 契约载荷 /
    find_ijob_script / call_ijob_brief_page / _validate_envelope / to_legacy_row）；
  - **不读真实工作区**：每个用例用 pytest `tmp_path` 现场造一个最小工作区
    （openpyxl 写 xlsx + 写巡检 md + 写 内容状态.json），可离线自包含运行；
  - CLI 端到端断言一律显式 `--workspace <tmp>`，绝不触碰真实总表。
"""
from __future__ import annotations

import json
import os
import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest

SCRIPTS = Path(__file__).resolve().parents[1] / "scripts"
sys.path.insert(0, str(SCRIPTS))

import emit_briefing as emit  # noqa: E402
import workspace_source as wsrc  # noqa: E402

SCRIPT = SCRIPTS / "emit_briefing.py"

# ══════════════════════════════════════════════════════════════════════
# fixtures / 构造器
# ══════════════════════════════════════════════════════════════════════

BATCH_HEADER = [
    "单位名称", "招聘类别", "招聘标题", "地区", "发布日期", "截止日期",
    "网申入口", "官方公告链接", "招聘人数", "学历要求", "年龄限制",
    "英语要求", "计算机要求", "咨询电话", "备注",
    "内容状态", "是否已发布",
]
COMPANY_HEADER = ["名称", "全称", "性质", "城市", "官网", "招聘官网", "所属行业"]

UNIT = "测试单位"
TITLE = "测试单位2027校招"
KEY = wsrc.state_key(UNIT, TITLE)


def batch_row(**overrides):
    """一条最小可用的 `招聘信息` 台账行。"""
    row = {
        "单位名称": UNIT, "招聘类别": "校园招聘", "招聘标题": TITLE,
        "地区": "南京", "发布日期": "2026-09-01", "截止日期": "2026-09-20",
        "网申入口": "https://apply.example/apply",
        "官方公告链接": "https://apply.example/notice",
        "招聘人数": "10人", "学历要求": "本科及以上", "年龄限制": "",
        "英语要求": "", "计算机要求": "", "咨询电话": "025-00000000",
        "备注": "", "内容状态": "unwritten", "是否已发布": "否",
    }
    row.update(overrides)
    return row


def company_row(**overrides):
    row = {
        "名称": UNIT, "全称": "测试单位有限公司", "性质": "国有企业",
        "城市": "南京市", "官网": "https://example.test",
        "招聘官网": "https://hr.example.test", "所属行业": "能源",
    }
    row.update(overrides)
    return row


def rec(unit=UNIT, title=TITLE, state="unwritten", facts=None, files=None):
    """`内容状态.json` 里的一个批次记录。"""
    return {
        "unit": unit, "title": title, "state": state, "history": [],
        "facts": facts or {}, "files": files or {}, "notes": "",
    }


def _write_xlsx(path, batch_rows=(), company_rows=(), batch_header=BATCH_HEADER,
                company_header=COMPANY_HEADER, drop_batch_cols=()):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = wsrc.SHEET_BATCHES
    header = [c for c in batch_header if c not in drop_batch_cols]
    ws.append(header)
    for r in batch_rows:
        ws.append([r.get(c, "") for c in header])
    ws2 = wb.create_sheet(wsrc.SHEET_COMPANIES)
    ws2.append(list(company_header))
    for r in company_rows:
        ws2.append([r.get(c, "") for c in company_header])
    wb.save(str(path))


def make_workspace(base, *, batch_rows=(), company_rows=(), state_batches=None,
                   reports=None, pieces=None, state_raw=None, make_state=True,
                   **xlsx_kw):
    """在 ``base/码上职业`` 造一个最小工作区，返回其根目录 Path。"""
    ws = base / "码上职业"
    ws.mkdir(parents=True, exist_ok=True)
    _write_xlsx(ws / wsrc.TOTAL_XLSX, batch_rows, company_rows, **xlsx_kw)
    if state_raw is not None:
        (ws / wsrc.STATE_JSON).write_text(state_raw, encoding="utf-8")
    elif make_state:
        write_state(ws, state_batches or {})
    for name, text in (reports or {}).items():
        d = ws / wsrc.PATROL_DIR
        d.mkdir(parents=True, exist_ok=True)
        (d / name).write_text(text, encoding="utf-8")
    for rel, names in (pieces or {}).items():
        d = ws / rel
        d.mkdir(parents=True, exist_ok=True)
        for n in names:
            (d / n).write_text("占位稿件", encoding="utf-8")
    return ws


def write_state(ws, batches):
    (ws / wsrc.STATE_JSON).write_text(
        json.dumps({"schema_version": "1.0.0", "states": list(wsrc.STATES),
                    "batches": batches}, ensure_ascii=False, indent=2),
        encoding="utf-8")


def read_state(ws):
    return json.loads((ws / wsrc.STATE_JSON).read_text(encoding="utf-8"))


# ── 巡检报告构造器 ──────────────────────────────────────────────────

def prow(date="2026-09-28", deadline="2026-10-10", layer="有效",
         source="机关单位来源", title="甲单位2027年公开招聘公告",
         url="https://a.example/notice"):
    return {"date": date, "deadline": deadline, "layer": layer,
            "source": source, "title": title, "url": url}


def patrol_md(units, ncol=5, tail=True):
    """units = [(单位名, 源URL, [prow, ...]), ...]；tail 追加 三之二~三之五 段。"""
    L = ["# 巡检报告", "", "## 一、概览", "", "共 1 条。", "",
         "## 三、当前在架招聘公告（真·开启·有效性核验后，按单位）", ""]
    for unit, unit_url, rows in units:
        L += [f"**{unit}**（{unit_url}）", ""]
        if ncol == 5:
            L += ["| 日期 | 报名截止 | 状态 | 来源 | 标题 |", "|---|---|---|---|---|"]
            for r in rows:
                L.append(f"| {r['date']} | {r['deadline']} | {r['layer']} | "
                         f"{r['source']} | [{r['title']}]({r['url']}) |")
        else:
            L += ["| 日期 | 报名截止 | 状态 | 标题 |", "|---|---|---|---|"]
            for r in rows:
                L.append(f"| {r['date']} | {r['deadline']} | {r['layer']} | "
                         f"[{r['title']}]({r['url']}) |")
        L.append("")
    if tail:
        L += [
            "## 三之二、招聘流程后续/变更（按单位汇总，不逐条列示）", "",
            "**三之二单位**（https://tail.example/）", "",
            "| 日期 | 报名截止 | 状态 | 来源 | 标题 |", "|---|---|---|---|---|",
            "| 2026-09-01 | 2026-10-01 | 有效 | 尾部来源 | "
            "[三之二不应计入](https://tail.example/a) |", "",
            "## 三之三、招聘会·引才活动", "",
            "| 日期 | 报名截止 | 状态 | 来源 | 标题 |", "|---|---|---|---|---|",
            "| 2026-09-02 | 2026-10-02 | 有效 | 尾部来源 | "
            "[三之三不应计入](https://tail.example/b) |", "",
            "## 三之四、招聘栏目入口（导航 tab，非公告，故不计入主口径）", "",
            "## 三之五、已过期/待核（不混入主口径，供核对与复盘）", "",
            "| 日期 | 报名截止 | 状态 | 来源 | 标题 |", "|---|---|---|---|---|",
            "| 2026-09-03 | 2026-10-03 | 待核 | 尾部来源 | "
            "[三之五不应计入](https://tail.example/c) |", "",
        ]
    return "\n".join(L) + "\n"


# ── legacy row 构造器（reconcile_draftbox / build_briefing 用） ──────

def legacy_row(cid=1, **overrides):
    row = {k: "" for k in wsrc.LEGACY_KEYS}
    row.update({
        "cid": cid, "id": cid, "html_len": 0, "rh": "",
        "content_state": "unwritten", "title": TITLE, "company": UNIT,
        "full_name": "测试单位有限公司", "batch_type": "校园招聘",
        "city": "南京", "publish_date": "2026-09-01", "deadline": "2026-09-20",
        "min_degree": "本科及以上", "total_vacancies": "10人",
        "apply_url": "https://apply.example/apply",
        "announcement_url": "https://apply.example/notice",
    })
    row.update(overrides)
    return row


def run_cli(ws, *args, env=None):
    return subprocess.run(
        [sys.executable, str(SCRIPT), "--workspace", str(ws), *args],
        capture_output=True, text=True, timeout=180, env=env)


# ══════════════════════════════════════════════════════════════════════
# 1. discover_workspace —— 工作区发现（显式/环境变量是硬约束，不回退）
# ══════════════════════════════════════════════════════════════════════

def test_discover_workspace_explicit_path_is_valid(tmp_path):
    ws = make_workspace(tmp_path, batch_rows=[batch_row()])
    assert wsrc.discover_workspace(str(ws)) == ws
    assert wsrc.discover_workspace(ws) == ws          # Path 也接受


def test_discover_workspace_explicit_invalid_raises_and_never_falls_back(tmp_path, monkeypatch):
    good = make_workspace(tmp_path / "good", batch_rows=[batch_row()])
    bad = tmp_path / "bad"
    bad.mkdir()
    # 环境变量指向合法工作区：显式非法时**必须报错，不得静默回退到它**
    monkeypatch.setenv(wsrc.WORKSPACE_ENV, str(good))
    with pytest.raises(wsrc.SourceError) as ei:
        wsrc.discover_workspace(str(bad))
    assert ei.value.code == "workspace_invalid"
    assert wsrc.TOTAL_XLSX in ei.value.message


def test_discover_workspace_env_valid(tmp_path, monkeypatch):
    ws = make_workspace(tmp_path, batch_rows=[batch_row()])
    monkeypatch.setenv(wsrc.WORKSPACE_ENV, str(ws))
    assert wsrc.discover_workspace() == ws


def test_discover_workspace_env_invalid_raises(tmp_path, monkeypatch):
    bad = tmp_path / "env-bad"
    bad.mkdir()
    monkeypatch.setenv(wsrc.WORKSPACE_ENV, str(bad))
    with pytest.raises(wsrc.SourceError) as ei:
        wsrc.discover_workspace()
    assert ei.value.code == "workspace_invalid"
    assert wsrc.WORKSPACE_ENV in ei.value.message


def test_discover_workspace_not_found(tmp_path, monkeypatch):
    monkeypatch.delenv(wsrc.WORKSPACE_ENV, raising=False)
    monkeypatch.setattr(wsrc, "DEFAULT_WORKSPACE", str(tmp_path / "无此工作区"))
    with pytest.raises(wsrc.SourceError) as ei:
        wsrc.discover_workspace()
    assert ei.value.code == "workspace_not_found"
    assert ei.value.details.get("checked") == [str(tmp_path / "无此工作区")]


# ══════════════════════════════════════════════════════════════════════
# 2. load_state —— 内容状态.json（唯一真源）
# ══════════════════════════════════════════════════════════════════════

def test_load_state_missing_raises_state_missing(tmp_path):
    ws = make_workspace(tmp_path, batch_rows=[batch_row()], make_state=False)
    with pytest.raises(wsrc.SourceError) as ei:
        wsrc.load_state(ws)
    assert ei.value.code == "state_missing"
    assert "state-init" in ei.value.details.get("hint", "")


def test_load_state_optional_returns_empty_state(tmp_path):
    ws = make_workspace(tmp_path, batch_rows=[batch_row()], make_state=False)
    state = wsrc.load_state(ws, required=False)
    assert state["batches"] == {}
    assert state["states"] == list(wsrc.STATES)
    assert state["schema_version"]


def test_load_state_bad_json_raises_state_invalid(tmp_path):
    ws = make_workspace(tmp_path, batch_rows=[batch_row()], state_raw="{不是 JSON,")
    for required in (True, False):   # 坏 JSON 即便 required=False 也必须报错
        with pytest.raises(wsrc.SourceError) as ei:
            wsrc.load_state(ws, required=required)
        assert ei.value.code == "state_invalid"


def test_load_state_missing_batches_dict_raises_state_invalid(tmp_path):
    ws = make_workspace(tmp_path, batch_rows=[batch_row()],
                        state_raw=json.dumps({"schema_version": "1.0.0"}))
    with pytest.raises(wsrc.SourceError) as ei:
        wsrc.load_state(ws)
    assert ei.value.code == "state_invalid"
    assert "batches" in ei.value.message

    # batches 不是字典也不行
    ws2 = make_workspace(tmp_path / "list", batch_rows=[batch_row()],
                         state_raw=json.dumps({"batches": []}))
    with pytest.raises(wsrc.SourceError) as ei2:
        wsrc.load_state(ws2)
    assert ei2.value.code == "state_invalid"


# ══════════════════════════════════════════════════════════════════════
# 3. load_batches —— xlsx 台账 → legacy row
# ══════════════════════════════════════════════════════════════════════

def test_load_batches_maps_xlsx_columns_to_legacy_keys(tmp_path):
    ws = make_workspace(tmp_path, batch_rows=[batch_row()])
    rows = wsrc.load_batches(ws, wsrc.load_state(ws))
    assert len(rows) == 1
    r = rows[0]
    assert r["company"] == UNIT
    assert r["batch_type"] == "校园招聘"
    assert r["title"] == TITLE
    assert r["city"] == "南京"
    assert r["publish_date"] == "2026-09-01"
    assert r["deadline"] == "2026-09-20"
    assert r["apply_url"] == "https://apply.example/apply"
    assert r["announcement_url"] == "https://apply.example/notice"
    assert r["source_url"] == r["announcement_url"]
    assert r["total_vacancies"] == "10人"
    assert r["min_degree"] == "本科及以上"
    assert r["consulting_phone"] == "025-00000000"
    assert r["cid"] == 1 and r["id"] == 1
    assert r["state_key"] == KEY
    # 台账没有的列一律空串（判缺口交给下游 grade，不允许 None 漏到下游）
    for k in ("full_name", "nature", "official_site", "employment_nature",
              "target_graduates", "wechat_url", "verdict_summary"):
        assert r[k] == ""
    assert set(wsrc.LEGACY_KEYS) <= set(r)


def test_state_json_overrides_xlsx_legacy_state_columns(tmp_path):
    ws = make_workspace(
        tmp_path, batch_rows=[batch_row(内容状态="unwritten", 是否已发布="否")],
        state_batches={KEY: rec(state="draft")})
    r = wsrc.load_batches(ws, wsrc.load_state(ws))[0]
    assert r["content_state"] == "draft"          # json 是真源
    assert r["xlsx_content_state"] == "unwritten"  # 总表旧列只作审计对照
    assert r["xlsx_published"] == "否"
    assert r["state_known"] is True


def test_content_state_falls_back_to_xlsx_when_json_has_no_record(tmp_path):
    ws = make_workspace(tmp_path, batch_rows=[batch_row(内容状态="written", 是否已发布="否")],
                        state_batches={})
    r = wsrc.load_batches(ws, wsrc.load_state(ws))[0]
    assert r["content_state"] == "written"
    assert r["state_known"] is False


def test_company_sheet_fills_unit_profile(tmp_path):
    ws = make_workspace(tmp_path, batch_rows=[batch_row()], company_rows=[company_row()])
    r = wsrc.load_batches(ws, wsrc.load_state(ws))[0]
    assert r["full_name"] == "测试单位有限公司"
    assert r["nature"] == "国有企业"
    assert r["headquarters_city"] == "南京市"
    assert r["official_site"] == "https://example.test"
    assert r["recruit_site"] == "https://hr.example.test"
    assert r["industry_catalog"] == "能源"
    assert r["city"] == "南京"                     # 台账地区不被公司档案覆盖


def test_facts_bag_fills_missing_batch_fields(tmp_path):
    ws = make_workspace(
        tmp_path,
        batch_rows=[batch_row(招聘人数="", 学历要求="")],
        state_batches={KEY: rec(state="written", facts={
            "employment_nature": "合同制", "target_graduates": "2027届",
            "line": "B", "verdict_summary": "官方公告已核",
            "hc": "99人", "edu": "硕士及以上",
        })})
    r = wsrc.load_batches(ws, wsrc.load_state(ws))[0]
    assert r["employment_nature"] == "合同制"
    assert r["target_graduates"] == "2027届"
    assert r["line"] == "B"
    assert r["verdict_summary"] == "官方公告已核"
    assert r["total_vacancies"] == "99人"     # facts.hc 仅在台账为空时补
    assert r["min_degree"] == "硕士及以上"     # facts.edu 同理

    # 台账已有值时 facts 不覆盖（台账是建稿当时登记值）
    ws2 = make_workspace(tmp_path / "b", batch_rows=[batch_row()],
                         state_batches={KEY: rec(facts={"hc": "99人", "edu": "硕士及以上"})})
    r2 = wsrc.load_batches(ws2, wsrc.load_state(ws2))[0]
    assert r2["total_vacancies"] == "10人"
    assert r2["min_degree"] == "本科及以上"


def test_missing_content_state_column_raises_xlsx_schema_changed(tmp_path):
    ws = make_workspace(tmp_path, batch_rows=[batch_row()],
                        drop_batch_cols=("内容状态",))
    with pytest.raises(wsrc.SourceError) as ei:
        wsrc.load_batches(ws, wsrc.load_state(ws))
    assert ei.value.code == "xlsx_schema_changed"
    assert "内容状态" in ei.value.message


def test_missing_published_column_raises_xlsx_schema_changed(tmp_path):
    ws = make_workspace(tmp_path, batch_rows=[batch_row()],
                        drop_batch_cols=("是否已发布",))
    with pytest.raises(wsrc.SourceError) as ei:
        wsrc.load_batches(ws, wsrc.load_state(ws))
    assert ei.value.code == "xlsx_schema_changed"
    assert "是否已发布" in ei.value.message


def test_article_text_is_read_only_when_local_piece_exists(tmp_path):
    ws = make_workspace(
        tmp_path, batch_rows=[batch_row()],
        state_batches={KEY: rec(state="written", files={"source": "2027年校园招聘/待发布/测试单位-源稿.md"})})
    rows = wsrc.load_batches(ws, wsrc.load_state(ws))
    assert rows[0]["rh"] == "" and rows[0]["article_length"] == 0   # 文件不存在 → 空

    p = ws / "2027年校园招聘/待发布/测试单位-源稿.md"
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text("正文" * 30, encoding="utf-8")
    r = wsrc.load_batches(ws, wsrc.load_state(ws))[0]
    assert r["rh"].startswith("正文")
    assert r["article_length"] == 60
    assert r["html_len"] == 60


# ══════════════════════════════════════════════════════════════════════
# 4. 巡检报告 —— parse_patrol_report / load_patrol / patrol_for
# ══════════════════════════════════════════════════════════════════════

def test_parse_patrol_report_five_column_table(tmp_path):
    p = tmp_path / "2026-09-30-机关单位招聘监测.md"
    p.write_text(patrol_md([("甲单位", "https://a.example/", [
        prow(layer="有效", source="南京市人社局",
             title="甲单位2027年公开招聘公告", url="https://a.example/n1")])]), encoding="utf-8")
    entries = wsrc.parse_patrol_report(p)
    assert len(entries) == 1
    e = entries[0]
    assert e["unit"] == "甲单位"
    assert e["date"] == "2026-09-28"
    assert e["deadline"] == "2026-10-10"
    assert e["layer"] == "有效"
    assert e["source"] == "南京市人社局"
    assert e["title"] == "甲单位2027年公开招聘公告"
    assert e["url"] == "https://a.example/n1"
    assert e["report"] == "2026-09-30-机关单位招聘监测.md"
    assert e["patrol_line"] == "机关单位"


def test_parse_patrol_report_four_column_table(tmp_path):
    p = tmp_path / "2026-09-30-银行招聘监测.md"
    p.write_text(patrol_md([("某银行", "https://b.example/", [
        prow(layer="长期", deadline="—", title="某银行2027校招",
             url="https://b.example/n1")])], ncol=4), encoding="utf-8")
    entries = wsrc.parse_patrol_report(p)
    assert len(entries) == 1
    assert entries[0]["layer"] == "长期"
    assert entries[0]["deadline"] == "—"
    assert entries[0]["source"] == ""            # 4 列表无「来源」
    assert entries[0]["url"] == "https://b.example/n1"
    assert entries[0]["patrol_line"] == "银行"


def test_parse_patrol_report_only_reads_section_three(tmp_path):
    p = tmp_path / "2026-09-30-机关单位招聘监测.md"
    p.write_text(patrol_md([("甲单位", "https://a.example/", [
        prow(title="主口径公告", url="https://a.example/keep")])]), encoding="utf-8")
    entries = wsrc.parse_patrol_report(p)
    urls = {e["url"] for e in entries}
    assert urls == {"https://a.example/keep"}
    for banned in ("https://tail.example/a", "https://tail.example/b", "https://tail.example/c"):
        assert banned not in urls
    assert not any(e["unit"] == "三之二单位" for e in entries)


def test_parse_patrol_report_filters_unknown_layer_status(tmp_path):
    p = tmp_path / "2026-09-30-铁路招聘监测.md"
    rows = [
        prow(date="2026-09-01", layer="有效", url="https://c.example/1"),
        prow(date="2026-09-02", layer="长期", url="https://c.example/2"),
        prow(date="2026-09-03", layer="待核", url="https://c.example/3"),
        prow(date="2026-09-04", layer="已过期", url="https://c.example/4"),
        prow(date="2026-09-05", layer="无效", url="https://c.example/5"),
    ]
    p.write_text(patrol_md([("丙单位", "https://c.example/", rows)]), encoding="utf-8")
    entries = wsrc.parse_patrol_report(p)
    assert [e["layer"] for e in entries] == ["有效", "长期", "待核"]
    assert {e["url"] for e in entries} == {f"https://c.example/{i}" for i in (1, 2, 3)}


def test_parse_patrol_report_unreadable_path_returns_empty(tmp_path):
    assert wsrc.parse_patrol_report(tmp_path / "不存在.md") == []
    plain = tmp_path / "随笔.md"
    plain.write_text("没有 ## 三、 段", encoding="utf-8")
    assert wsrc.parse_patrol_report(plain) == []


def test_load_patrol_reads_recent_monitor_reports(tmp_path):
    mk = lambda unit, url: patrol_md([(unit, "https://a.example/", [prow(url=url)])])
    ws = make_workspace(tmp_path, batch_rows=[batch_row()], reports={
        "2026-09-28-机关单位招聘监测.md": mk("甲单位", "https://a.example/d28"),
        "2026-09-29-机关单位招聘监测.md": mk("甲单位", "https://a.example/d29"),
        "2026-09-30-机关单位招聘监测.md": mk("甲单位", "https://a.example/d30"),
        "随笔.md": "不是巡检报告",
    })
    entries, reports = wsrc.load_patrol(ws)
    assert reports == ["2026-09-30-机关单位招聘监测.md",
                       "2026-09-29-机关单位招聘监测.md",
                       "2026-09-28-机关单位招聘监测.md"]
    assert len(entries) == 3

    entries2, reports2 = wsrc.load_patrol(ws, max_reports=2)
    assert reports2 == reports[:2]
    assert len(entries2) == 2
    assert {e["url"] for e in entries2} == {"https://a.example/d30", "https://a.example/d29"}


def test_load_patrol_without_patrol_dir_returns_empty(tmp_path):
    ws = make_workspace(tmp_path, batch_rows=[batch_row()])
    assert wsrc.load_patrol(ws) == ([], [])


def _parse_two(ws, older, newer, name_a="2026-09-28-机关单位招聘监测.md",
               name_b="2026-09-30-机关单位招聘监测.md"):
    """把两份报告写进同一工作区并返回解析出的条目。"""
    d = ws / wsrc.PATROL_DIR
    d.mkdir(parents=True, exist_ok=True)
    (d / name_a).write_text(older, encoding="utf-8")
    (d / name_b).write_text(newer, encoding="utf-8")
    return wsrc.load_patrol(ws)[0]


def test_patrol_for_dedupes_same_url_preferring_harder_layer(tmp_path):
    ws = make_workspace(tmp_path, batch_rows=[batch_row()])
    same = "https://a.example/notice"
    entries = _parse_two(
        ws,
        patrol_md([("甲单位", "https://a.example/", [prow(layer="长期", url=same)])]),
        patrol_md([("甲单位", "https://a.example/", [prow(layer="有效", url=same)])]))
    hits = wsrc.patrol_for("甲单位", entries)
    assert len(hits) == 1                      # 同一 URL 跨报告去重
    assert hits[0]["layer"] == "有效"           # 分层更硬者胜
    assert hits[0]["report"].startswith("2026-09-30")


def test_patrol_for_prefers_newer_report_within_same_layer(tmp_path):
    ws = make_workspace(tmp_path, batch_rows=[batch_row()])
    same = "https://a.example/notice"
    entries = _parse_two(
        ws,
        patrol_md([("甲单位", "https://a.example/", [prow(layer="有效", url=same)])]),
        patrol_md([("甲单位", "https://a.example/", [prow(layer="有效", url=same)])]))
    hits = wsrc.patrol_for("甲单位", entries)
    assert len(hits) == 1
    assert hits[0]["report"].startswith("2026-09-30")


def test_patrol_for_layer_beats_report_recency(tmp_path):
    ws = make_workspace(tmp_path, batch_rows=[batch_row()])
    same = "https://a.example/notice"
    # 旧报告=有效，新报告=长期 → 「有效」优先，即使它来自更旧的报告
    entries = _parse_two(
        ws,
        patrol_md([("甲单位", "https://a.example/", [prow(layer="有效", url=same)])]),
        patrol_md([("甲单位", "https://a.example/", [prow(layer="长期", url=same)])]))
    hits = wsrc.patrol_for("甲单位", entries)
    assert len(hits) == 1
    assert hits[0]["layer"] == "有效"
    assert hits[0]["report"].startswith("2026-09-28")


def test_patrol_for_matches_unit_title_and_source(tmp_path):
    ws = make_workspace(tmp_path, batch_rows=[batch_row()])
    text = patrol_md([
        ("甲单位", "https://a.example/", [prow(title="甲单位2027年公开招聘公告",
                                            source="南京市人社局", url="https://a.example/1")]),
    ])
    entries = _parse_two(ws, patrol_md([]), text)
    assert len(entries) == 1
    assert wsrc.patrol_for("甲单位", entries)[0]["url"] == "https://a.example/1"
    assert wsrc.patrol_for("南京市人社局", entries)[0]["url"] == "https://a.example/1"
    assert wsrc.patrol_for("2027年公开招聘", entries)[0]["url"] == "https://a.example/1"
    assert wsrc.patrol_for("丁单位", entries) == []
    assert wsrc.patrol_for("", entries) == []
    assert wsrc.patrol_for("  ", entries) == []


def test_patrol_for_respects_limit_and_sorts_by_layer(tmp_path):
    ws = make_workspace(tmp_path, batch_rows=[batch_row()])
    text = patrol_md([("甲单位", "https://a.example/", [
        prow(layer="待核", url="https://a.example/wait"),
        prow(layer="长期", url="https://a.example/long"),
        prow(layer="有效", url="https://a.example/ok"),
    ])])
    entries = _parse_two(ws, patrol_md([]), text)
    hits = wsrc.patrol_for("甲单位", entries)
    assert [e["layer"] for e in hits] == ["有效", "长期", "待核"]
    assert len(wsrc.patrol_for("甲单位", entries, limit=2)) == 2
    assert len(wsrc.patrol_for("甲单位", entries, limit=1)) == 1


def test_patrol_for_dedupes_by_title_when_url_missing(tmp_path):
    ws = make_workspace(tmp_path, batch_rows=[batch_row()])
    d = ws / wsrc.PATROL_DIR
    d.mkdir(parents=True, exist_ok=True)
    body = "\n".join([
        "## 三、当前在架招聘公告", "",
        "**甲单位**（https://a.example/）", "",
        "| 日期 | 报名截止 | 状态 | 来源 | 标题 |", "|---|---|---|---|---|",
        "| 2026-09-28 | — | 有效 | 来源 | 无链接公告 |", ""])
    (d / "2026-09-28-机关单位招聘监测.md").write_text(body, encoding="utf-8")
    (d / "2026-09-30-机关单位招聘监测.md").write_text(body, encoding="utf-8")
    entries = wsrc.load_patrol(ws)[0]
    assert len(entries) == 2 and all(e["url"] == "" for e in entries)
    hits = wsrc.patrol_for("甲单位", entries)
    assert len(hits) == 1 and hits[0]["report"].startswith("2026-09-30")


# ══════════════════════════════════════════════════════════════════════
# 5. init_state —— 从总表播种 + 目录补充；幂等 / force
# ══════════════════════════════════════════════════════════════════════

def test_init_state_seeds_from_xlsx_and_scan_dirs(tmp_path):
    ws = make_workspace(
        tmp_path,
        batch_rows=[batch_row(内容状态="written", 是否已发布="否")],
        pieces={"2027年校园招聘/待发布": ["某单位2027校招-源稿.md",
                                          "某单位2027校招-排版.html"]},
        make_state=False)
    state, info = wsrc.init_state(ws)
    assert info["action"] == "created"
    assert info["seeded"] == 1
    assert info["from_dirs"] == 1
    assert info["batches"] == 2
    assert (ws / wsrc.STATE_JSON).is_file()

    seeded = state["batches"][KEY]
    assert seeded["state"] == "written"           # 取自总表「内容状态」
    assert seeded["history"][0]["to"] == "written"
    assert "初始化" in seeded["history"][0]["note"]

    dir_key = [k for k in state["batches"] if "某单位2027校招" in k]
    assert len(dir_key) == 1
    assert state["batches"][dir_key[0]]["state"] == "written"
    assert state["batches"][dir_key[0]]["notes"] == "台账缺登记"


def test_init_state_kept_when_file_exists(tmp_path):
    ws = make_workspace(tmp_path, batch_rows=[batch_row(内容状态="unwritten")],
                        state_batches={KEY: rec(state="published")})
    state, info = wsrc.init_state(ws)
    assert info["action"] == "kept"
    assert info["batches"] == 1
    assert state["batches"][KEY]["state"] == "published"   # 不覆盖既有真源


def test_init_state_force_rebuilds_from_xlsx(tmp_path):
    ws = make_workspace(tmp_path, batch_rows=[batch_row(内容状态="written")],
                        state_batches={KEY: rec(state="published")})
    state, info = wsrc.init_state(ws, force=True)
    assert info["action"] == "rebuilt"
    assert info["seeded"] == 1
    assert state["batches"][KEY]["state"] == "written"     # 回落到总表值


def test_init_state_records_published_from_xlsx(tmp_path):
    ws = make_workspace(tmp_path, batch_rows=[batch_row(内容状态="published", 是否已发布="是")],
                        make_state=False)
    state, _info = wsrc.init_state(ws)
    assert state["batches"][KEY]["state"] == "published"


def test_init_state_strips_piece_filename_prefixes(tmp_path):
    ws = make_workspace(
        tmp_path, batch_rows=[batch_row()],
        pieces={"教育系统/待发布": [
            "2026-09-30-南京特殊教育师范学院第二批专任教师5人-源稿.md",
            "02-04-中国核电-源稿.md"]},
        make_state=False)
    state, info = wsrc.init_state(ws)
    assert info["from_dirs"] == 2
    units = {v["unit"] for v in state["batches"].values()
             if v["notes"] == "台账缺登记"}
    assert units == {"南京特殊教育师范学院第二批专任教师5人", "中国核电"}
    # 只剥「日期/期号」前缀，批次描述保留
    assert wsrc._unit_from_piece("2026-09-30-南师大第二批") == "南师大第二批"
    assert wsrc._unit_from_piece("02-04-中国核电") == "中国核电"
    assert wsrc._unit_from_piece("中国核电") == "中国核电"


def test_init_state_ignores_non_piece_files(tmp_path):
    ws = make_workspace(
        tmp_path, batch_rows=[batch_row()],
        pieces={"2027年校园招聘/待发布": ["某单位-源稿.md", "随笔.md", "某单位-封面.png"]},
        make_state=False)
    _state, info = wsrc.init_state(ws)
    assert info["from_dirs"] == 1          # 非三件套后缀的文件不建记录


# ══════════════════════════════════════════════════════════════════════
# 6. set_batch_state —— 状态推进 + history 留痕
# ══════════════════════════════════════════════════════════════════════

def test_set_batch_state_advances_and_appends_history(tmp_path):
    ws = make_workspace(tmp_path, batch_rows=[batch_row()])
    r1 = wsrc.set_batch_state(ws, UNIT, TITLE, "written", by="AI", note="初稿")
    assert r1["state"] == "written"
    assert r1["history"][-1]["from"] == "unwritten"
    assert r1["history"][-1]["to"] == "written"
    assert r1["history"][-1]["by"] == "AI"

    r2 = wsrc.set_batch_state(ws, UNIT, TITLE, "draft", by="AI")
    assert r2["history"][-1]["from"] == "written"
    assert r2["history"][-1]["to"] == "draft"
    assert len(r2["history"]) == 2                 # 追加，不覆盖

    persisted = wsrc.load_state(ws)["batches"][KEY]
    assert persisted["state"] == "draft"
    assert len(persisted["history"]) == 2


def test_set_batch_state_same_state_does_not_duplicate_history(tmp_path):
    ws = make_workspace(tmp_path, batch_rows=[batch_row()])
    wsrc.set_batch_state(ws, UNIT, TITLE, "written")
    r = wsrc.set_batch_state(ws, UNIT, TITLE, "written")
    assert r["state"] == "written"
    assert len(r["history"]) == 1


def test_set_batch_state_creates_record_for_unknown_batch(tmp_path):
    ws = make_workspace(tmp_path, batch_rows=[batch_row()])
    r = wsrc.set_batch_state(ws, "台账外单位", "", "written", note="--mark 新建")
    assert r["state"] == "written"
    assert r["history"][-1]["from"] == "unwritten"
    assert r["history"][-1]["to"] == "written"


def test_set_batch_state_unknown_state_raises(tmp_path):
    ws = make_workspace(tmp_path, batch_rows=[batch_row()])
    with pytest.raises(wsrc.SourceError) as ei:
        wsrc.set_batch_state(ws, UNIT, TITLE, "done")
    assert ei.value.code == "state_unknown"
    assert "unwritten" in ei.value.message
    # 报错时不落盘
    assert wsrc.load_state(ws)["batches"] == {}


def test_set_batch_state_records_files(tmp_path):
    ws = make_workspace(tmp_path, batch_rows=[batch_row()])
    r = wsrc.set_batch_state(ws, UNIT, TITLE, "written",
                             files={"source": "2027年校园招聘/待发布/测试单位-源稿.md",
                                    "cover": ""})
    assert r["files"] == {"source": "2027年校园招聘/待发布/测试单位-源稿.md"}


# ══════════════════════════════════════════════════════════════════════
# 7. audit —— 四方对账（只读）
# ══════════════════════════════════════════════════════════════════════

def test_audit_lists_state_xlsx_mismatch(tmp_path):
    ws = make_workspace(tmp_path, batch_rows=[batch_row(内容状态="unwritten", 是否已发布="否")],
                        state_batches={KEY: rec(state="written")})
    rep = wsrc.audit(ws)
    assert rep["batches_in_state"] == 1
    assert rep["batches_in_xlsx"] == 1
    assert len(rep["state_xlsx_mismatch"]) == 1
    mm = rep["state_xlsx_mismatch"][0]
    assert mm["unit"] == UNIT
    assert mm["json"] == "written"
    assert mm["xlsx"] == "unwritten"
    assert rep["xlsx_rows_missing_in_state"] == []
    assert rep["state_counts"] == {"written": 1}


def test_audit_reports_xlsx_rows_missing_in_state(tmp_path):
    ws = make_workspace(tmp_path, batch_rows=[batch_row(内容状态="unwritten")],
                        state_batches={})
    rep = wsrc.audit(ws)
    assert len(rep["xlsx_rows_missing_in_state"]) == 1
    assert rep["xlsx_rows_missing_in_state"][0]["unit"] == UNIT
    assert rep["xlsx_counts"] == {"unwritten": 1}


def test_audit_flags_published_without_local_archive(tmp_path):
    ws = make_workspace(tmp_path, batch_rows=[batch_row(内容状态="published", 是否已发布="是")],
                        state_batches={KEY: rec(state="published")})
    rep = wsrc.audit(ws)
    assert rep["published_not_archived"] == [UNIT]
    assert rep["local_archived"]["2027年校园招聘/已发布"] == 0

    # 归档件到位（红线 11）→ 该条不再出现
    ws2 = make_workspace(
        tmp_path / "archived",
        batch_rows=[batch_row(内容状态="published", 是否已发布="是")],
        state_batches={KEY: rec(state="published")},
        pieces={"2027年校园招聘/已发布": ["测试单位2027校招-源稿.md"]})
    rep2 = wsrc.audit(ws2)
    assert rep2["published_not_archived"] == []
    assert rep2["local_archived"]["2027年校园招聘/已发布"] == 1


def test_audit_counts_local_inventory_per_line(tmp_path):
    ws = make_workspace(
        tmp_path, batch_rows=[batch_row()], state_batches={KEY: rec()},
        pieces={"教育系统/待发布": ["A单位-源稿.md", "B单位-源稿.md"],
                "事业单位/待发布": ["C单位-源稿.md"]})
    rep = wsrc.audit(ws)
    assert rep["local_inventory"]["教育系统/待发布"] == 2
    assert rep["local_inventory"]["事业单位/待发布"] == 1
    assert rep["state_file"] == str(ws / wsrc.STATE_JSON)


def test_audit_is_read_only(tmp_path):
    ws = make_workspace(tmp_path, batch_rows=[batch_row()], state_batches={KEY: rec()})
    before = (ws / wsrc.STATE_JSON).read_bytes()
    wsrc.audit(ws)
    assert (ws / wsrc.STATE_JSON).read_bytes() == before


# ══════════════════════════════════════════════════════════════════════
# 8. push_state_to_xlsx —— 单向回推（⛔ 只允许 tmp 工作区）
# ══════════════════════════════════════════════════════════════════════

@pytest.fixture
def iso_backup_dir(tmp_path, monkeypatch):
    """把 push 的备份目录也锁进 tmp_path，测试结束后不残留 /tmp 文件。"""
    monkeypatch.setattr(wsrc.tempfile, "gettempdir", lambda: str(tmp_path))
    return tmp_path


def _xlsx_cells(path):
    wb = openpyxl.load_workbook(path)
    try:
        sh = wb[wsrc.SHEET_BATCHES]
        header = [c.value for c in sh[1]]
        return [dict(zip(header, (sh.cell(row=i, column=j + 1).value
                                  for j in range(len(header)))))
                for i in range(2, sh.max_row + 1)]
    finally:
        wb.close()


def test_push_state_to_xlsx_noop_when_already_synced(tmp_path, iso_backup_dir):
    ws = make_workspace(tmp_path, batch_rows=[batch_row(内容状态="written", 是否已发布="否")],
                        state_batches={KEY: rec(state="written")})
    res = wsrc.push_state_to_xlsx(ws)
    assert res["updated"] == 0
    assert Path(res["backup"]).is_file()
    assert len(res["sha256"]) == 64
    assert res["rollback"].startswith("cp '")
    row = _xlsx_cells(ws / wsrc.TOTAL_XLSX)[0]
    assert row["内容状态"] == "written"
    assert row["是否已发布"] == "否"


def test_push_state_to_xlsx_updates_both_state_columns(tmp_path, iso_backup_dir):
    ws = make_workspace(tmp_path, batch_rows=[batch_row(内容状态="unwritten", 是否已发布="否")],
                        state_batches={KEY: rec(state="published")})
    res = wsrc.push_state_to_xlsx(ws)
    assert res["updated"] == 1
    assert Path(res["backup"]).is_file()
    assert len(res["sha256"]) == 64
    assert f"'{res['backup']}'" in res["rollback"]
    assert "单位与招聘信息总表.xlsx" in res["backup"]

    row = _xlsx_cells(ws / wsrc.TOTAL_XLSX)[0]
    assert row["内容状态"] == "published"
    assert row["是否已发布"] == "是"

    # 再推一次 → 已同步，不再更新
    assert wsrc.push_state_to_xlsx(ws)["updated"] == 0


def test_push_state_to_xlsx_maps_draft_to_written_not_published(tmp_path, iso_backup_dir):
    ws = make_workspace(tmp_path, batch_rows=[batch_row(内容状态="unwritten", 是否已发布="否")],
                        state_batches={KEY: rec(state="draft")})
    assert wsrc.push_state_to_xlsx(ws)["updated"] == 1
    row = _xlsx_cells(ws / wsrc.TOTAL_XLSX)[0]
    assert row["内容状态"] == "written"     # 总表无 draft 档
    assert row["是否已发布"] == "否"


def test_push_state_to_xlsx_skips_batches_not_in_state(tmp_path, iso_backup_dir):
    ws = make_workspace(tmp_path, batch_rows=[batch_row(内容状态="written", 是否已发布="否")],
                        state_batches={})
    assert wsrc.push_state_to_xlsx(ws)["updated"] == 0
    row = _xlsx_cells(ws / wsrc.TOTAL_XLSX)[0]
    assert row["内容状态"] == "written"


# ══════════════════════════════════════════════════════════════════════
# 9. emit_briefing 纯函数 —— draft_state / grade* / url 归一
# ══════════════════════════════════════════════════════════════════════

@pytest.mark.parametrize(("state", "expected"), [
    ("unwritten", "库内未排"),
    ("written", "库内已成稿"),
    ("draft", "已推草稿箱"),
    ("published", "已发布"),
])
def test_draft_state_maps_all_four_states(state, expected):
    assert emit.draft_state({"content_state": state}) == expected


def test_draft_state_fallback_branch_never_infers_draftbox():
    # 无 content_state（或未知值）→ 回退到「可发布链接 + 正文长度」自证
    assert emit.draft_state({}) == "库内未排"
    assert emit.draft_state({"content_state": ""}) == "库内未排"
    assert emit.draft_state({"content_state": "bogus"}) == "库内未排"
    assert emit.draft_state({"content_state": "bogus", "html_len": 51}) == "库内已成稿"
    assert emit.draft_state({"content_state": "", "article_length": 51}) == "库内已成稿"
    assert emit.draft_state({"content_state": "", "html_len": 50}) == "库内未排"
    assert emit.draft_state({"content_state": "", "wechat_url": "short"}) == "库内未排"
    assert emit.draft_state(
        {"content_state": "", "wechat_url": "https://mp.weixin.qq.com/s/aaaa"}) == "已发布"
    # 本地有正文 ≠ 已推草稿箱
    assert emit.draft_state({"content_state": "written", "html_len": 99999}) == "库内已成稿"


@pytest.mark.parametrize(("value", "expected"), [
    ("", "EMPTY"),
    (None, "EMPTY"),
    ("   ", "EMPTY"),
    ("2026-09-20", "OK"),
    ("2026-09-01", "OK"),
    ("10月15日", "OK"),
    ("本科及以上", "OK"),
    ("10人", "OK"),
    ("待核", "VAGUE"),
    ("待定", "VAGUE"),
    ("预计10月", "VAGUE"),
    ("暂无", "VAGUE"),
    ("滚动招聘", "VAGUE"),
    ("招满即止", "VAGUE"),
    ("详见公告", "VAGUE"),
    ("以公告为准", "VAGUE"),
    ("—", "VAGUE"),
    ("--", "VAGUE"),
])
def test_grade_confidence_levels(value, expected):
    assert emit.grade(value) == expected


def test_deadline_is_derived_placeholder():
    assert emit.deadline_is_derived_placeholder(
        {"publish_date": "2026-09-01", "deadline": "2026-10-01"})            # +30 天
    assert not emit.deadline_is_derived_placeholder(
        {"publish_date": "2026-09-01", "deadline": "2026-09-25"})            # +24 天
    # 空值 / 非日期都不能判成占位
    assert not emit.deadline_is_derived_placeholder({"publish_date": "", "deadline": "2026-10-01"})
    assert not emit.deadline_is_derived_placeholder({"publish_date": "2026-09-01", "deadline": ""})
    assert not emit.deadline_is_derived_placeholder(
        {"publish_date": "2026-09-01", "deadline": "招满即止"})


def test_grade_deadline_downgrades_placeholder_and_pending_wording():
    base = {"publish_date": "2026-09-01", "deadline": "2026-09-25",
            "batch_type": "", "title": "", "verdict_summary": ""}

    def row(**kw):
        r = dict(base)
        r.update(kw)
        return r

    assert emit.grade_deadline(row()) == "OK"
    assert emit.grade_deadline(row(deadline="")) == "EMPTY"
    assert emit.grade_deadline(row(deadline=emit.PLACEHOLDER_DEADLINE)) == "VAGUE"
    assert emit.grade_deadline(row(deadline="2026-10-01")) == "VAGUE"      # 发布日 +30 天
    assert emit.grade_deadline(row(deadline="2026-11-30", batch_type="待核批次")) == "VAGUE"
    assert emit.grade_deadline(row(deadline="2026-11-30", title="截止待官方公告")) == "VAGUE"
    assert emit.grade_deadline(row(deadline="2026-11-30", verdict_summary="通告未披露")) == "VAGUE"
    assert emit.grade_deadline(row(deadline="招满即止")) == "VAGUE"
    # 真实日期 + 30 天以外的差值保持 OK
    assert emit.grade_deadline(row(deadline="2027-08-31")) == "OK"


@pytest.mark.parametrize(("raw", "expected"), [
    ("https://x.example/a", ("https://x.example/a", "")),
    ("rsj.taizhou.gov.cn（通知公告栏）", ("rsj.taizhou.gov.cn", "通知公告栏")),
    ("https://x.example/a (备注)", ("https://x.example/a", "备注")),
    ("https://x.example/a", ("https://x.example/a", "")),
])
def test_clean_url_splits_url_and_note(raw, expected):
    assert emit.clean_url(raw) == expected


@pytest.mark.parametrize(("raw", "expected"), [
    ("https://Example.COM/apply/", "example.com/apply"),
    ("  https://x.example/a//  ", "x.example/a"),
    ("", ""),
    # ⚠️ 当前实现先「去协议前缀」再 lower()，故大写协议去不掉（已在测试报告登记，未改生产码）
    ("HTTP://X.example/a（⛔ 51job专题）", "http://x.example/a"),
])
def test_norm_apply_url(raw, expected):
    assert emit.norm_apply_url(raw) == expected


def test_grade_gates_and_gate_keys():
    row = {k: "" for k in wsrc.LEGACY_KEYS}
    assert emit.grade_gates(row) == "EMPTY"
    row["cet_requirement"] = "CET-4 425"
    assert emit.grade_gates(row) == "OK"
    assert set(emit.GATE_KEYS) <= set(wsrc.LEGACY_KEYS)


# ══════════════════════════════════════════════════════════════════════
# 10. reconcile_draftbox / filter_rows（本地台账驱动，仍然成立的部分）
# ══════════════════════════════════════════════════════════════════════

def _rows():
    return [
        legacy_row(cid=11, title="南京大学2027校招", company="南京大学",
                   full_name="南京大学", apply_url="https://nju.example"),
        legacy_row(cid=12, title="中远海运科技2027校招", company="中远海运科技",
                   full_name="中远海运科技有限公司", apply_url="https://cosco.example"),
        legacy_row(cid=13, title="国网江苏2027校招", company="国网江苏",
                   full_name="国网江苏省电力有限公司", apply_url="https://sgcc.example",
                   content_state="written", article_length=100, html_len=100,
                   rh="正文" + "国网江苏2027校招" + "x" * 80),
    ]


def test_reconcile_exact_match_and_body_fallback():
    matched, unmatched, desync, semantic = emit.reconcile_draftbox(_rows(), [
        "南京大学2027校招",
        "完全不存在的单位",
    ])
    assert len(matched) == 1
    assert matched[0][1] == 11
    assert matched[0][4] == "标题精确"
    assert unmatched == ["完全不存在的单位"]
    assert semantic == []
    # 精确命中 + 本地状态=库内未排 → ⛔确定失同步（这正是对账的意义所在）
    assert desync == [(11, "南京大学", "南京大学2027校招")]

    # 草稿标题与批次标题毫无字面关系 → 靠稿件正文（rh）命中
    rows = _rows()
    rows[2] = legacy_row(cid=13, title="国网江苏电力2027校园招聘正式启动",
                         content_state="written", article_length=100, html_len=100,
                         rh="前文" + "某某神秘标题甲乙丙丁" + "后文" * 40)
    matched, _u, _d, _s = emit.reconcile_draftbox(rows, ["某某神秘标题甲乙丙丁"])
    assert matched[0][1] == 13
    assert matched[0][4] == "正文含标题"


def test_reconcile_body_match_ignores_rows_without_article():
    rows = _rows()
    rows[2] = legacy_row(cid=13, title="国网江苏电力2027校园招聘正式启动",
                         content_state="unwritten", article_length=0, html_len=0, rh="")
    _m, unmatched, _d, _s = emit.reconcile_draftbox(rows, ["某某神秘标题甲乙丙丁"])
    assert "某某神秘标题甲乙丙丁" in unmatched


def test_reconcile_duplicate_title_is_ambiguous_candidate():
    rows = _rows()
    rows.append(legacy_row(cid=14, title="南京大学2027校招", company="南京大学",
                           full_name="南京大学"))
    matched, _u, _d, semantic = emit.reconcile_draftbox(rows, ["南京大学2027校招"])
    assert len(matched) == 1
    assert matched[0][1] == 11
    assert matched[0][4] == "标题多批次候选:11,14"
    assert semantic == [matched[0]]
    assert emit._is_ambiguous_match(matched[0][4])
    best = emit._best_draft_matches(matched)
    assert set(best) == {11, 14}


def test_reconcile_alias_and_subsidiary_guard():
    rows = _rows()
    rows[1].update(company="华为南研所", full_name="华为技术有限公司南京研究所")
    rows[1]["title"] = "华为南研所2027届校园招聘"
    matched, unmatched, desync, semantic = emit.reconcile_draftbox(rows, ["华为南研所校招来了"])
    assert unmatched == []
    assert matched[0][1] == 12
    assert matched[0][4].startswith("别名候选:")
    assert semantic == [matched[0]]
    assert desync == []

    # 通用词不跨城市/跨机构误配
    rows2 = _rows()
    rows2[1].update(company="江苏省血液中心", full_name="江苏省血液中心")
    rows2[1]["title"] = "江苏省血液中心公开招聘"
    matched2, unmatched2, _d2, _s2 = emit.reconcile_draftbox(rows2, [
        "南京市红十字血液中心招聘",
        "南京市疾病预防控制中心招聘",
    ])
    assert matched2 == []
    assert len(unmatched2) == 2


def test_match_marker_and_strength_ordering():
    assert emit.match_marker("库内未排", "主体多批次候选:1,2") == "  ⚠️仅候选"
    assert emit.match_marker("库内未排", "标题精确") == "  ⛔确定失同步"
    assert emit.match_marker("库内已成稿", "标题精确") == ""
    exact = ("某草稿", 9, "单位", "库内未排", "标题精确")
    body = ("某草稿", 9, "单位", "库内未排", "正文含标题")
    alias = ("某草稿", 9, "单位", "库内未排", "别名候选:单位")
    ambiguous = ("某草稿", 9, "单位", "库内未排", "标题多批次候选:9,10")
    assert emit._best_draft_matches([alias, ambiguous, body, exact])[9] == exact
    assert emit._best_draft_matches([alias, ambiguous, body])[9] == body


def test_filter_rows_sorts_and_filters_by_keyword():
    rows = [
        legacy_row(cid=3, deadline="2026-10-01", company="测试单位", full_name="测试单位"),
        legacy_row(cid=1, deadline="2026-09-01", company="别的单位", full_name="别的单位",
                   title="别的单位2027校招"),
        legacy_row(cid=2, deadline="2026-09-15", company="测试单位",
                   full_name="测试单位", title="二次批次"),
    ]
    out = emit.filter_rows(rows, "测试单位")
    assert [r["cid"] for r in out] == [2, 3]          # 命中单位，按截止日升序
    assert emit.filter_rows(rows, "不存在") == []
    assert [r["cid"] for r in emit.filter_rows(rows)] == [1, 2, 3]   # 不传关键词=全部
    # 关键词也命中 地区/批次/行业（不止单位名）
    assert [r["cid"] for r in emit.filter_rows(rows, "南京")] == [1, 2, 3]


# ══════════════════════════════════════════════════════════════════════
# 11. build_briefing —— 投料单文案（数据源口径 + 巡检条目）
# ══════════════════════════════════════════════════════════════════════

def test_build_briefing_points_at_workspace_sources_not_external_project():
    text = emit.build_briefing(legacy_row(), "2026-09-30 09:00")
    assert "# BRIEFING · 测试单位有限公司" in text
    assert "本工作区台账" in text
    assert "内容状态.json" in text
    assert "巡检报告" in text
    assert "ijob" not in text.lower()
    for banned in ("jobs.db", "campaign_articles", "src/ijob"):
        assert banned not in text
    assert "**本文件不是核实结论，是核实起点**" in text


def test_build_briefing_blocks_on_vague_deadline_and_lists_gap():
    text = emit.build_briefing(legacy_row(deadline="2026-10-01", publish_date="2026-09-01"),
                               "2026-09-30 09:00")
    assert "暂不可动笔" in text
    assert "报名截止" in text
    assert "回写 `内容状态.json`" in text


def test_build_briefing_renders_patrol_entries():
    patrol = [{"layer": "有效", "deadline": "2026-10-10",
               "title": "甲单位2027年公开招聘公告",
               "url": "https://a.example/notice",
               "source": "南京市人社局", "report": "2026-09-30-机关单位招聘监测.md"}]
    text = emit.build_briefing(legacy_row(), "2026-09-30 09:00", patrol)
    assert "[甲单位2027年公开招聘公告](https://a.example/notice)" in text
    assert "2026-09-30-机关单位招聘监测.md" in text
    assert "## C2. 巡检报告·在架条目" in text


def test_build_briefing_says_report_gap_is_not_absence():
    text = emit.build_briefing(legacy_row(), "2026-09-30 09:00", [])
    assert "巡检报告里没匹配到" in text
    assert "不等于没在招" in text


# ══════════════════════════════════════════════════════════════════════
# 12. 源码卫生 —— 彻底删除 ijob 依赖（回归闸）
# ══════════════════════════════════════════════════════════════════════

def test_production_scripts_have_no_external_project_remnants():
    for name in ("emit_briefing.py", "workspace_source.py"):
        src = (SCRIPTS / name).read_text(encoding="utf-8").lower()
        assert "ijob" not in src, f"{name} 仍出现 ijob"
        for banned in ("sqlite3", "jobs.db", "ijob_skill_script"):
            assert banned not in src, f"{name} 仍出现 {banned}"


def test_removed_adapter_symbols_are_gone():
    for gone in ("to_legacy_row", "find_ijob_script", "call_ijob_brief_page",
                 "_validate_envelope", "fetch_brief_items", "IjobContractError",
                 "IJOB_SKILL_SCRIPT", "IJOB_SCRIPT_ENV"):
        assert not hasattr(emit, gone), f"emit_briefing 仍暴露 {gone}"
    assert not hasattr(wsrc, "to_legacy_row")
    assert hasattr(emit, "wsrc") and emit.wsrc is wsrc


def test_tests_depend_only_on_local_production_modules():
    """测试只依赖本技能的两个生产模块 + 标准库；不得再引入外部数据适配模块。"""
    import ast

    tree = ast.parse(Path(__file__).read_text(encoding="utf-8"))
    mods = set()
    for node in ast.walk(tree):
        if isinstance(node, ast.Import):
            mods.update(a.name.split(".")[0] for a in node.names)
        elif isinstance(node, ast.ImportFrom) and node.module:
            mods.add(node.module.split(".")[0])
    assert {"emit_briefing", "workspace_source"} <= mods
    assert mods <= {"__future__", "ast", "json", "os", "subprocess", "sys",
                    "pathlib", "openpyxl", "pytest",
                    "emit_briefing", "workspace_source"}


# ══════════════════════════════════════════════════════════════════════
# 13. CLI 端到端（全部指向 tmp 工作区）
# ══════════════════════════════════════════════════════════════════════

def test_cli_state_init_then_audit(tmp_path):
    ws = make_workspace(tmp_path, batch_rows=[batch_row(内容状态="unwritten")],
                        make_state=False)

    p1 = run_cli(ws, "--state-init")
    assert p1.returncode == 0, p1.stderr
    assert "[ok] created" in p1.stdout
    assert "seeded: 1" in p1.stdout
    assert (ws / wsrc.STATE_JSON).is_file()

    p2 = run_cli(ws, "--state-init")
    assert p2.returncode == 0
    assert "[ok] kept" in p2.stdout        # 幂等

    p3 = run_cli(ws, "--state-init", "--force")
    assert p3.returncode == 0
    assert "[ok] rebuilt" in p3.stdout

    p4 = run_cli(ws, "--state-audit")
    assert p4.returncode == 0, p4.stderr
    assert "内容状态对账" in p4.stdout
    assert "本报告只读" in p4.stdout
    assert str(ws / wsrc.STATE_JSON) in p4.stdout


def test_cli_gap_only_lists_missing_fields(tmp_path):
    ws = make_workspace(tmp_path, batch_rows=[batch_row(截止日期="待核", 学历要求="")],
                        state_batches={KEY: rec()})
    p = run_cli(ws, UNIT, "--gap-only")
    assert p.returncode == 0, p.stderr
    assert UNIT in p.stdout
    assert "报名截止" in p.stdout
    assert "学历要求" in p.stdout
    assert "[红线未登记]" in p.stdout


def test_cli_mark_advances_state_and_records_facts(tmp_path):
    ws = make_workspace(tmp_path, batch_rows=[batch_row()], state_batches={})
    p = run_cli(ws, "--mark", UNIT, "--to", "draft",
                "--facts", "employment_nature=合同制")
    assert p.returncode == 0, p.stderr
    assert "[ok] 测试单位 ｜" in p.stdout
    assert "draft" in p.stdout
    saved = read_state(ws)["batches"][KEY]
    assert saved["state"] == "draft"
    assert saved["facts"]["employment_nature"] == "合同制"
    assert saved["history"][-1]["to"] == "draft"

    # --mark 不带 --to → 明确报错，不许静默成功
    p2 = run_cli(ws, "--mark", UNIT)
    assert p2.returncode != 0
    assert "--to" in (p2.stdout + p2.stderr)

    # --facts 形式错误 → facts_invalid
    p3 = run_cli(ws, "--mark", UNIT, "--to", "written", "--facts", "缺等号")
    assert p3.returncode != 0
    assert "facts_invalid" in (p3.stdout + p3.stderr)


def test_cli_mark_unknown_keyword_is_not_silent(tmp_path):
    ws = make_workspace(tmp_path, batch_rows=[batch_row()], state_batches={})
    p = run_cli(ws, "--mark", "台账里没有的单位", "--to", "written")
    # --mark 允许本地留痕直接建（不是错误），但必须打印新建提示
    assert p.returncode == 0, p.stderr
    assert "新建" in p.stdout
    assert "台账里没有的单位" in read_state(ws)["batches"][wsrc.state_key("台账里没有的单位", "")]["unit"]


def test_cli_ready_loose_works_offline_without_draftbox(tmp_path):
    ws = make_workspace(
        tmp_path,
        batch_rows=[batch_row(截止日期="2099-12-31")],
        company_rows=[company_row()],
        state_batches={KEY: rec(facts={"employment_nature": "合同制",
                                       "target_graduates": "2027届"})})
    p = run_cli(ws, "--ready-loose", "--no-draftbox")
    assert p.returncode == 0, p.stderr
    assert "宽松档" in p.stdout
    assert "测试单位有限公司" in p.stdout
    assert "不可直接采信" in p.stdout        # 草稿箱未对账 → 名单需人工核
    assert "用工性质" in p.stdout            # 宽松档必须显式提示仍缺的必填项


def test_cli_explicit_invalid_workspace_never_falls_back_to_real_workspace(tmp_path):
    bad = tmp_path / "空目录"
    bad.mkdir()
    env = dict(os.environ)
    env.pop(wsrc.WORKSPACE_ENV, None)

    # 非状态子命令
    p1 = subprocess.run(
        [sys.executable, str(SCRIPT), "--workspace", str(bad), UNIT, "--gap-only"],
        capture_output=True, text=True, timeout=180, env=env)
    assert p1.returncode == 1
    assert "workspace_invalid" in p1.stderr
    assert "不得把失败当成" in p1.stderr
    assert "测试单位" not in p1.stdout

    # 状态子命令同样拒绝
    p2 = subprocess.run(
        [sys.executable, str(SCRIPT), "--workspace", str(bad), "--state-init"],
        capture_output=True, text=True, timeout=180, env=env)
    assert p2.returncode == 1
    assert "workspace_invalid" in p2.stderr
    assert not (bad / wsrc.STATE_JSON).exists()


def test_cli_missing_state_file_reports_state_missing(tmp_path):
    ws = make_workspace(tmp_path, batch_rows=[batch_row()], make_state=False)
    p = run_cli(ws, UNIT, "--gap-only")
    assert p.returncode == 1
    assert "state_missing" in p.stderr
    assert "state-init" in p.stderr


# ══════════════════════════════════════════════════════════════════════
# 14. import_draft_titles —— 草稿箱标题导入（本次改造同批新增）
# ══════════════════════════════════════════════════════════════════════

def test_import_draft_titles_dry_run_writes_nothing(tmp_path):
    ws = make_workspace(tmp_path, batch_rows=[batch_row()], state_batches={KEY: rec()})
    before = (ws / wsrc.STATE_JSON).read_bytes()
    res = wsrc.import_draft_titles(ws, ["某某新单位2027校招：面向全省招聘"])
    assert res["applied"] == 0
    assert res["planned"] == ["某某新单位2027校招：面向全省招聘"]
    assert res["skipped"] == []
    assert (ws / wsrc.STATE_JSON).read_bytes() == before     # 预演不落盘


def test_import_draft_titles_never_downgrades_existing_state(tmp_path):
    ws = make_workspace(tmp_path, batch_rows=[batch_row(内容状态="published", 是否已发布="是")],
                        state_batches={KEY: rec(state="published")})
    res = wsrc.import_draft_titles(ws, ["测试单位2027校招（二次修改版）"], apply=True)
    assert res["planned"] == []
    assert len(res["skipped"]) == 1
    assert res["skipped"][0][1] == UNIT
    assert res["skipped"][0][2] == "published"
    # ⛔ 草稿箱里的旧草稿不得把 published 拉回 draft
    assert wsrc.load_state(ws)["batches"][KEY]["state"] == "published"


def test_import_draft_titles_apply_creates_draft_records(tmp_path):
    ws = make_workspace(tmp_path, batch_rows=[batch_row()], state_batches={KEY: rec()})
    title = "某某新单位2027校招：面向全省招聘"
    res = wsrc.import_draft_titles(ws, [title], apply=True, by="AI")
    assert res["applied"] == 1
    assert res["planned"] == [title]

    saved = wsrc.load_state(ws)["batches"]
    key = wsrc.state_key("某某新单位2027校招", title)
    assert saved[key]["state"] == "draft"
    assert saved[key]["unit"] == "某某新单位2027校招"
    assert saved[key]["history"][-1]["to"] == "draft"
    assert "草稿箱" in saved[key]["notes"]
    assert len(saved) == 2                                   # 原批次未被改写


def test_import_draft_titles_empty_list_is_a_noop(tmp_path):
    ws = make_workspace(tmp_path, batch_rows=[batch_row()], state_batches={KEY: rec()})
    before = (ws / wsrc.STATE_JSON).read_bytes()
    res = wsrc.import_draft_titles(ws, [], apply=True)
    assert res == {"planned": [], "skipped": [], "applied": 0}
    assert (ws / wsrc.STATE_JSON).read_bytes() == before


def test_unit_from_title_strips_bracket_prefix_and_tail():
    assert wsrc._unit_from_title("【原创】某某单位2027校招") == "某某单位2027校招"
    assert wsrc._unit_from_title("某某单位2027校招：公告") == "某某单位2027校招"
    assert wsrc._unit_from_title("短") == "短"
