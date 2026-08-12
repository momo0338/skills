#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
把 dy_wode.py 输出的 dy_wode.json 转成多工作表 Excel（dy_wode.xlsx）。

每个分类一个工作表（喜欢/收藏/观看历史/稍后再看/消息），另加「汇总」页。

用法：
  python3 export_xlsx.py [dy_wode.json 路径]
  缺省参数时读取当前目录下的 dy_wode.json，输出同目录 dy_wode.xlsx。
"""

import json
import os
import sys
from collections import Counter


def main() -> int:
    src = sys.argv[1] if len(sys.argv) > 1 else "dy_wode.json"
    src = os.path.abspath(src)
    if not os.path.exists(src):
        sys.exit(f"未找到 {src}，请先运行 dy_wode.py 生成 JSON。")
    rows = json.load(open(src, encoding="utf-8"))
    if not rows:
        sys.exit("JSON 为空。")

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        sys.exit("需要 openpyxl：pip install openpyxl")

    cols = list(rows[0].keys())
    by_cat: dict[str, list[dict]] = {}
    for r in rows:
        by_cat.setdefault(r.get("分类") or "未分类", []).append(r)

    wb = Workbook()
    wb.remove(wb.active)
    header_fill = PatternFill("solid", fgColor="2F5B3A")
    header_font = Font(color="FFFFFF", bold=True)
    stat_cols = {"点赞数", "评论数", "收藏数", "转发数"}

    # 汇总页
    ws = wb.create_sheet("汇总")
    ws.append(["分类", "条数"])
    total = 0
    for cat, items in sorted(by_cat.items(), key=lambda kv: -len(kv[1])):
        ws.append([cat, len(items)])
        total += len(items)
    ws.append(["合计", total])
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 10
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font

    for cat, items in by_cat.items():
        sheet = wb.create_sheet(cat)
        sheet.append(cols)
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
        for r in items:
            sheet.append([r.get(c, "") for c in cols])
        for i, c in enumerate(cols, 1):
            letter = get_column_letter(i)
            if c in stat_cols:
                width = 10
            elif c == "视频名称":
                width = 60
            elif c in ("视频ID",):
                width = 24
            else:
                width = 14
            sheet.column_dimensions[letter].width = width
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

    out = os.path.join(os.path.dirname(src), "dy_wode.xlsx")
    wb.save(out)
    print(f"已导出：{out}（{total} 条，{len(by_cat)} 个分类工作表）")
    return 0


if __name__ == "__main__":
    sys.exit(main())
